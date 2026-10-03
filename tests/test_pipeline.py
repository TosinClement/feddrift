"""FedDrift tests: critical transformations, label honesty, packaging restrictions and release-gate conditions.

Run: python -m pytest -q tests
Tests that need generated outputs (data/processed/...) skip with a reason if those outputs are absent.
"""
import importlib.util
import json
import os
import shutil
import subprocess
import sys

import pandas as pd
import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
CODE = os.path.join(REPO, "code")
sys.path.insert(0, CODE)


def load(name):
    spec = importlib.util.spec_from_file_location(name.replace(".py", "").replace("-", "_"), os.path.join(CODE, name))
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


def need(path):
    p = os.path.join(REPO, path)
    if not os.path.exists(p):
        pytest.skip(f"{path} not generated yet")
    return p


def idx(n, start="2000-01-01"):
    return pd.date_range(start, periods=n, freq="MS").strftime("%Y-%m-%d")


# ---------------------------------------------------------------- transformations
def test_pair_stats_rebase_depth_and_censoring():
    m = load("04_vintage_pairs.py")
    a = pd.Series(range(100, 136), index=idx(36), dtype=float)
    r = m.pair_stats("X", "2003-01-15", "2003-02-15", a, a * 2.0)
    assert r["n_revised"] == 36 and r["rebase_like"] and r["revision_depth_months"] == 37 and r["depth_censored"]
    c = a.copy()
    c.iloc[-2:] += 1.0                                     # Nov and Dec 2002 revised; vintage Feb 2003
    r2 = m.pair_stats("X", "2003-01-15", "2003-02-15", a, c)
    assert r2["n_revised"] == 2 and not r2["rebase_like"] and r2["revision_depth_months"] == 3
    assert not r2["depth_censored"]
    b = a.copy()
    b.loc["2002-12-01"] = 999                              # unchanged values elsewhere
    b = pd.concat([b, pd.Series([140.0], index=["2003-01-01"])])
    r3 = m.pair_stats("X", "2003-01-15", "2003-02-15", a, b)
    assert r3["n_new_obs"] == 1 and r3["n_revised"] == 1


def test_no_revision_is_not_a_drift_event():
    m = load("04_vintage_pairs.py")
    a = pd.Series([1.0, 2.0, 3.0], index=idx(3))
    b = pd.concat([a, pd.Series([4.0], index=["2000-04-01"])])
    r = m.pair_stats("X", "2000-04-15", "2000-05-15", a, b)
    assert r["n_revised"] == 0 and r["n_new_obs"] == 1 and "earliest_revised_obs" not in r


def test_vintage_hash_is_canonical():
    m = load("10_qa.py")
    h1 = m.vintage_hash(["2000-02-01", "2000-01-01"], ["2.0", "1.0"])
    h2 = m.vintage_hash(["2000-01-01", "2000-02-01"], ["1.0", "2.0"])
    assert h1 == h2 and h1 != m.vintage_hash(["2000-01-01", "2000-02-01"], ["1.00", "2.0"])


def test_rule_engine_routes_events():
    m = load("05_propose_labels.py")
    rules = pd.read_csv(os.path.join(REPO, "config", "revision_rules.csv"), dtype=str).fillna("")
    rules["priority"] = rules.priority.astype(int)
    rules = rules.sort_values("priority")
    base = dict(n_revised=2, n_dropped_obs=0, n_obs_next=10, n_obs_prev=9, rebase_like=False, depth_censored=False)
    cases = [
        (dict(series_id="DGORDER", vintage_month=6, revision_depth_months=3), "R23_ROUTINE_M3"),
        (dict(series_id="DGORDER", vintage_month=5, revision_depth_months=160), "R30_CENSUS_ANNUAL"),
        (dict(series_id="DGORDER", vintage_month=8, revision_depth_months=160), "R90_UNCLASSIFIED"),
        (dict(series_id="CPIAUCSL", vintage_month=2, revision_depth_months=61), "R10_BLS_SA_FEB"),
        (dict(series_id="CPIAUCSL", vintage_month=2, revision_depth_months=19, depth_censored=True), "R10_BLS_SA_FEB"),
        (dict(series_id="CPIAUCNS", vintage_month=2, revision_depth_months=61), "R90_UNCLASSIFIED"),
        (dict(series_id="TSIFRGHT", vintage_month=9, revision_depth_months=250), "R26_ROUTINE_TSI"),
        (dict(series_id="RSAFS", vintage_month=4, revision_depth_months=100, rebase_like=True), "R02_REBASE"),
        (dict(series_id="RSAFS", vintage_month=4, revision_depth_months=0, n_revised=0), "R00_EXTENSION"),
        (dict(series_id="CPIAUCSL", vintage_month=4, revision_depth_months=0, n_revised=0, n_dropped_obs=1,
              n_obs_next=19, n_obs_prev=19), "R01_ARCHIVE_WINDOW"),
    ]
    for over, expected in cases:
        row = pd.Series({**base, **over})
        assert m.propose(row, rules) == expected, (over, expected)


def test_event_ids_unique_and_stable_format():
    p = need("data/processed/proposed_events.csv")
    ev = pd.read_csv(p)
    assert ev.event_id.is_unique
    assert (ev.event_id == "FD-" + ev.series_id + "-" + ev.vintage_next.str.replace("-", "")).all()


def test_t2_split_assignment_is_leakage_free():
    m = load("09_benchmark.py")
    d = pd.DataFrame({"first_vintage": ["2003-01-15", "2010-02-15", "2012-06-15", "2016-03-15", "2014-01-15"],
                      "mature_vintage": ["2006-02-15", "2013-03-15", "2015-07-15", "2019-04-15", "2017-02-15"]})
    assert list(m.assign_t2_split(d)) == ["train", "validation", "validation", "test", "excluded_gap"]


# ---------------------------------------------------------------- label honesty
def _label_root(tmp_path):
    for d in ["data/processed", "labels/decisions", "evidence", "config"]:
        os.makedirs(tmp_path / d, exist_ok=True)
    shutil.copy(need("data/processed/proposed_events.csv"), tmp_path / "data/processed")
    for f in ["source_registry.csv", "rule_evidence.csv", "cluster_evidence.csv"]:
        shutil.copy(need(f"evidence/{f}"), tmp_path / "evidence")
    return tmp_path


def _apply(root):
    env = dict(os.environ, FEDDRIFT_ROOT=str(root))
    return subprocess.run([sys.executable, os.path.join(CODE, "08_apply_labels.py")], env=env,
                          capture_output=True, text=True)


def test_nothing_verified_without_decisions(tmp_path):
    root = _label_root(tmp_path)
    assert _apply(root).returncode == 0
    ev = pd.read_csv(root / "data/processed/drift_events.csv", keep_default_na=False)
    de = ev[ev.is_drift_event.astype(str) == "True"]
    assert set(de.label_status) <= {"rule_pending_review", "release_pending_review"}
    assert (de.final_cause == "").all() and (de.reviewer == "").all()
    st = json.load(open(root / "data/processed/label_status.json"))
    assert st["n_verified"] == 0 and st["n_unverified"] == st["n_drift_events"]


def test_decisions_apply_and_require_reviewer(tmp_path):
    root = _label_root(tmp_path)
    cl = pd.read_csv(root / "evidence/cluster_evidence.csv")
    c1, c2 = cl.release_cluster_id.iloc[0], cl.release_cluster_id.iloc[1]
    pd.DataFrame([{"rule_id": "R23_ROUTINE_M3", "decision": "approve", "verified_by": "Test Reviewer",
                   "verified_date": "2026-10-03", "notes": ""},
                  {"rule_id": "R25_ROUTINE_PPI", "decision": "approve", "verified_by": "",   # no reviewer: ignored
                   "verified_date": "2026-10-03", "notes": ""}]).to_csv(root / "labels/decisions/rule_decisions.csv", index=False)
    pd.DataFrame([{"release_cluster_id": c1, "decision": "unknown", "final_cause": "", "final_secondary_cause": "",
                   "verified_by": "Test Reviewer", "verified_date": "2026-10-03", "notes": ""},
                  {"release_cluster_id": c2, "decision": "set", "final_cause": "correction", "final_secondary_cause": "",
                   "verified_by": "Test Reviewer", "verified_date": "2026-10-03", "notes": ""}]
                 ).to_csv(root / "labels/decisions/release_decisions.csv", index=False)
    assert _apply(root).returncode == 0
    ev = pd.read_csv(root / "data/processed/drift_events.csv", keep_default_na=False)
    m3 = ev[ev.proposed_rule == "R23_ROUTINE_M3"]
    assert (m3.label_status == "rule_verified").all() and (m3.final_cause == "advance_to_revised").all()
    assert (m3.reviewer == "Test Reviewer").all() and (m3.label_source_urls != "").all()
    ppi = ev[ev.proposed_rule == "R25_ROUTINE_PPI"]
    assert (ppi.label_status == "rule_pending_review").all()
    a = ev[ev.release_cluster_id == c1]
    assert (a.label_status == "unknown_verified").all() and (a.final_cause == "unknown").all()
    b = ev[ev.release_cluster_id == c2]
    assert (b.label_status == "release_verified").all() and (b.final_cause == "correction").all()


def test_import_review_rejects_invalid_rows(tmp_path):
    from openpyxl import Workbook
    root = _label_root(tmp_path)
    shutil.copy(os.path.join(REPO, "config", "causes.csv"), root / "config")
    wb = Workbook()
    wb.active.title = "Judgment calls"
    wb["Judgment calls"].append(["decision_id", "decision", "verified_by", "verified_date", "notes"])
    for name, header, row in [
        ("Rules", ["rule_id", "decision", "verified_by", "verified_date", "notes"],
         ["R23_ROUTINE_M3", "approve", "", "2026-10-03", ""]),                     # missing reviewer
        ("Releases", ["release_cluster_id", "decision", "final_cause", "final_secondary_cause", "verified_by",
                      "verified_date", "notes"], ["FDC-X", "set", "made_up_cause", "", "T", "2026-10-03", ""]),
        ("Event overrides", ["event_id", "final_cause", "final_secondary_cause", "reason", "source_url", "verified_by",
                             "verified_date"], ["FD-NOPE-20200101", "correction", "", "", "", "T", "2026-10-03"])]:
        ws = wb.create_sheet(name)
        ws.append(header)
        ws.append(row)
    x = tmp_path / "r.xlsx"
    wb.save(x)
    env = dict(os.environ, FEDDRIFT_ROOT=str(root))
    r = subprocess.run([sys.executable, os.path.join(CODE, "import_review.py"), str(x)], env=env,
                       capture_output=True, text=True)
    assert r.returncode != 0
    assert "verified_by is empty" in r.stderr and "made_up_cause" in r.stderr and "unknown event_id" in r.stderr
    assert not os.path.exists(root / "labels/decisions/rule_decisions.csv")


# ---------------------------------------------------------------- packaging and gate
def test_package_refuses_alfred_values_and_keys(tmp_path):
    pkg = load("13_package.py")
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "leak.json").write_text('[{"realtime_start": "2020-01-01", "realtime_end": "2020-02-01", '
                                                 '"date": "2019-12-01", "value": "1.0"}]')
    (tmp_path / "data" / "leak.csv").write_text("realtime_start,realtime_end,date,value\n")
    (tmp_path / "notes.md").write_text("url?api_key=" + "a" * 32)
    (tmp_path / "data" / "raw").mkdir()
    (tmp_path / "data" / "raw" / "alfred_cache").mkdir()
    (tmp_path / "data" / "raw" / "alfred_cache" / "X.json").write_text("cached")
    old = pkg.ROOT
    pkg.ROOT = str(tmp_path)
    try:
        inc, probs = pkg.content_problems()
    finally:
        pkg.ROOT = old
    assert not any("alfred_cache" in f for f in inc)
    assert any("leak.json" in p for p in probs) and any("leak.csv" in p for p in probs) and any("notes.md" in p for p in probs)


def _gate_project(tmp_path, readme, labels_ok=True, final_mode=True):
    for d in ["config", "data/processed", "paper"]:
        os.makedirs(tmp_path / d, exist_ok=True)
    shutil.copy(os.path.join(REPO, "config", "placeholders.json"), tmp_path / "config")
    (tmp_path / "README.md").write_text(readme)
    json.dump({"n_unverified": 0 if labels_ok else 5}, open(tmp_path / "data/processed/label_status.json", "w"))
    json.dump({"checks": [{"id": "Q01", "status": "PASS"}]}, open(tmp_path / "data/processed/qa.json", "w"))
    json.dump({"final_mode": final_mode}, open(tmp_path / "paper/stats.json", "w"))
    return subprocess.run([sys.executable, os.path.join(CODE, "publish_gate.py"), str(tmp_path), "--stage", "final"],
                          capture_output=True, text=True), \
        subprocess.run([sys.executable, os.path.join(CODE, "publish_gate.py"), str(tmp_path), "--stage", "pre-deposit"],
                       capture_output=True, text=True)


def test_gate_blocks_drafts_placeholders_and_unverified(tmp_path):
    fin, pre = _gate_project(tmp_path / "a", "Clean text.")
    assert fin.returncode == 0 and pre.returncode == 0
    fin, pre = _gate_project(tmp_path / "b", "DOI {{ZENODO_DOI}}")
    assert fin.returncode == 1 and pre.returncode == 0
    fin, pre = _gate_project(tmp_path / "c", "Status: DRAFT")
    assert fin.returncode == 1 and pre.returncode == 1
    fin, pre = _gate_project(tmp_path / "d", "see {{SOMETHING_ELSE}}")
    assert pre.returncode == 1
    fin, pre = _gate_project(tmp_path / "e", "Clean.", labels_ok=False)
    assert fin.returncode == 1 and pre.returncode == 1
    fin, pre = _gate_project(tmp_path / "f", "Clean.", final_mode=False)
    assert fin.returncode == 1


def test_figures_final_mode_refused_while_labels_unverified():
    st = json.load(open(need("data/processed/label_status.json")))
    if st["n_unverified"] == 0:
        pytest.skip("all labels verified; refusal path not reachable")
    r = subprocess.run([sys.executable, os.path.join(CODE, "11_figures_stats.py"), "--final"], capture_output=True,
                       text=True, cwd=REPO)
    assert r.returncode != 0 and "refused" in (r.stderr + r.stdout)


def test_secret_scan_detects_tokens(tmp_path):
    qa = load("10_qa.py")
    (tmp_path / "x.txt").write_text("token " + "github_pat_" + "A" * 40)
    old = qa.ROOT
    qa.ROOT = str(tmp_path)
    try:
        hits = qa.secret_scan()
    finally:
        qa.ROOT = old
    assert hits and hits[0][1].startswith("GitHub")
