#!/usr/bin/env python3
"""07 - Build the label owner's review package.

Writes (machine-generated, safe to regenerate; contains NO author decisions):
  labels/review/rules_for_review.csv        routine rules with event counts, agency quotes, recommendation
  labels/review/releases_for_review.csv     release clusters with evidence, automated check, recommendation
  labels/review/judgment_calls.csv          project-level decisions with evidence and consequences
  labels/review/FedDrift_label_review.xlsx  the same, as one workbook with drop-down decision cells, plus the
                                            full event inventory and the source registry for reference

Author decisions live ONLY in labels/decisions/*.csv (written by code/import_review.py from the workbook, or by
hand). This script never reads or writes them, so regenerating the review package cannot overwrite a decision.
"""

import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from fd_common import CONFIG, PROCESSED, ROOT

REVIEW = os.path.join(ROOT, "labels", "review")
EVID = os.path.join(ROOT, "evidence")
FONT = Font(name="Arial", size=9)
BOLD = Font(name="Arial", size=9, bold=True)
HEAD_FILL = PatternFill("solid", start_color="DDE6F0")
INPUT_FILL = PatternFill("solid", start_color="FFFF00")
WRAP = Alignment(wrap_text=True, vertical="top")

JUDGMENT_CALLS = [
    {"decision_id": "J1", "topic": "Redistribution of ALFRED/FRED values",
     "evidence": "FRED Legal (fred.stlouisfed.org/legal): series tagged 'Public Domain: Citation Requested' 'may be used "
                 "without permission, provided you do not engage in any prohibited use' with citation 'Source: <agency> via "
                 "FRED'. All 17 panel series carry that tag (data/manifests/alfred_manifest.json). FRED API Terms of Use: "
                 "third-party series may need owner permission; prohibited to replicate the essential FRED experience; "
                 "applications must display 'This product uses the FRED® API but is not endorsed or certified by the "
                 "Federal Reserve Bank of St. Louis.'",
     "recommendation": "keep_no_values",
     "options": "keep_no_values: ship manifests + reconstruction code only (current; matches the original brief; robust "
                "to future tag changes; users need a free FRED key). | ship_values: also ship the 6,312 vintage vectors "
                "with 'Source: <agency> via FRED' citations (permitted by the tag as read on 2026-10-03; easier reuse; "
                "must re-check tags at every release and drop any series that changes tag)."},
    {"decision_id": "J2", "topic": "Derived revision statistics in the release table",
     "evidence": "Percentage magnitudes, growth revisions and KS statistics are computed from ALFRED values. Under J1 "
                 "the underlying series are public domain (citation requested); the statistics are summaries, not "
                 "observation levels, and do not replicate FRED.",
     "recommendation": "ship_derived",
     "options": "ship_derived: keep the columns in drift_events.csv (benchmark usable without a key; T1 features "
                "available to everyone). | reconstruct_only: drop the magnitude columns from the release and let users "
                "rebuild them with a key (stricter; T1 then requires a FRED key)."},
    {"decision_id": "J3", "topic": "Full BTS Monthly Transportation Statistics file",
     "evidence": "BTS labels the dataset Public Domain U.S. Government, but it compiles 136 columns from many providers; "
                 "the Socrata metadata carries no per-column source notes. FedDrift needs only the two BTS-produced TSI "
                 "columns.",
     "recommendation": "withhold_full_file",
     "options": "withhold_full_file: ship only the extracted TSI columns, plus URL and SHA-256 of the full file "
                "(current). | ship_full_file: include crem-w557.csv as fetched."},
    {"decision_id": "J4", "topic": "Release-level review granularity",
     "evidence": "442 non-routine events form 220 agency releases (same program, same vintage date). One release has "
                 "one agency document and one cause; its events differ only by series.",
     "recommendation": "review_per_release",
     "options": "review_per_release: one decision per release applies to all its listed events (label status "
                "'release_verified'); single-event exceptions via the Event overrides sheet. | review_per_event: one "
                "decision per event (442 rows)."},
    {"decision_id": "J5", "topic": "Evidence bar for a non-unknown label",
     "evidence": "Recommendations are graded strong / moderate / weak / none (see Releases sheet).",
     "recommendation": "strong_or_moderate",
     "options": "strong_or_moderate: approve a cause when a release-specific or program-level agency source supports "
                "it; set 'unknown' for weak/none unless you find a better source. | strong_only: only release-specific "
                "sources with a matching footprint; everything else becomes 'unknown'. | any: approve weak "
                "recommendations too (not recommended)."},
]


def style_sheet(ws, df, input_cols=(), widths=None, comments=None):
    ws.append(list(df.columns))
    for c in ws[1]:
        c.font, c.fill, c.alignment = BOLD, HEAD_FILL, WRAP
    for row in df.itertuples(index=False):
        ws.append([("" if pd.isna(v) else v) for v in row])
    cols = list(df.columns)
    for j, name in enumerate(cols, 1):
        letter = get_column_letter(j)
        ws.column_dimensions[letter].width = (widths or {}).get(name, 16)
        for i in range(2, ws.max_row + 1):
            cell = ws[f"{letter}{i}"]
            cell.font, cell.alignment = FONT, WRAP
            if name in input_cols:
                cell.fill = INPUT_FILL
        if comments and name in comments:
            ws[f"{letter}1"].comment = Comment(comments[name], "FedDrift")
    ws.freeze_panes = "B2"


def add_list(ws, df, col, values):
    j = list(df.columns).index(col) + 1
    letter = get_column_letter(j)
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{letter}2:{letter}{max(2, len(df) + 1)}")


def main():
    os.makedirs(REVIEW, exist_ok=True)
    ev = pd.read_csv(os.path.join(PROCESSED, "proposed_events.csv"))
    rules = pd.read_csv(os.path.join(CONFIG, "revision_rules.csv"), dtype=str).fillna("")
    causes = pd.read_csv(os.path.join(CONFIG, "causes.csv"))
    final_causes = list(causes[causes.allowed_as_final == "yes"].cause)
    reg = pd.read_csv(os.path.join(EVID, "source_registry.csv"), dtype=str).fillna("")
    rev = pd.read_csv(os.path.join(EVID, "rule_evidence.csv"), dtype=str).fillna("")
    cev = pd.read_csv(os.path.join(EVID, "cluster_evidence.csv"), dtype=str).fillna("")
    regi = reg.set_index("source_id")

    def src_text(ids):
        out = []
        for s in [x for x in ids.split(";") if x]:
            if s in regi.index:
                r = regi.loc[s]
                out.append(f"{s}: {r.publisher}, \"{r.title}\" ({r.publication_date or 'date n/a'}) {r.url}")
        return "\n".join(out)

    # ---- rules
    cnt = ev.groupby("proposed_rule").agg(n_events=("event_id", "size"), n_series=("series_id", "nunique"),
                                          first_vintage=("vintage_next", "min"), last_vintage=("vintage_next", "max"))
    rr = rules[rules.review_level == "rule"].copy()
    rr = rr.join(cnt, on="rule_id")
    rr["agency_sources"] = rr.rule_id.map(lambda r: src_text(";".join(rev[rev.rule_id == r].source_id)))
    rr["agency_quotes"] = rr.rule_id.map(lambda r: "\n---\n".join(q for q in rev[rev.rule_id == r].quote if q))
    rr["claude_recommendation"] = "approve"
    rr["decision"] = ""
    rr["verified_by"] = ""
    rr["verified_date"] = ""
    rr["notes"] = ""
    rr = rr[["rule_id", "proposed_cause", "series", "routine_window_months", "condition", "rationale", "n_events",
             "n_series", "first_vintage", "last_vintage", "agency_sources", "agency_quotes", "claude_recommendation",
             "decision", "verified_by", "verified_date", "notes"]]
    rr.to_csv(os.path.join(REVIEW, "rules_for_review.csv"), index=False)

    # ---- releases
    rel = cev.copy()
    rel["sources"] = rel.source_ids.map(src_text)
    rel["event_ids"] = rel.release_cluster_id.map(
        lambda c: ";".join(sorted(ev[ev.release_cluster_id == c].event_id)))
    rel["decision"] = ""
    rel["final_cause"] = ""
    rel["final_secondary_cause"] = ""
    rel["verified_by"] = ""
    rel["verified_date"] = ""
    rel["notes"] = ""
    conf_order = {"none": 0, "weak": 1, "moderate": 2, "strong": 3}
    rel = rel.sort_values(["release_program", "vintage_date"])
    rel = rel[["release_cluster_id", "release_program", "vintage_date", "series", "n_events", "event_ids",
               "proposed_rules", "earliest_revised_obs", "max_depth_months", "any_depth_censored", "sources",
               "evidence_locator", "evidence_quote", "agency_issue_date", "automated_check", "recommended_cause",
               "recommended_secondary_cause", "confidence", "recommendation_basis", "decision", "final_cause",
               "final_secondary_cause", "verified_by", "verified_date", "notes"]]
    rel.to_csv(os.path.join(REVIEW, "releases_for_review.csv"), index=False)

    jc = pd.DataFrame(JUDGMENT_CALLS)
    jc["decision"] = ""
    jc["verified_by"] = ""
    jc["verified_date"] = ""
    jc["notes"] = ""
    jc.to_csv(os.path.join(REVIEW, "judgment_calls.csv"), index=False)

    # ---- workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    lines = [
        ("FedDrift label review workbook", BOLD),
        ("Reviewer: Tosin Clement (label owner). Nothing in this workbook is a verified label until you fill the yellow cells.", FONT),
        ("Yellow cells are the only cells to edit. Everything else is machine-generated evidence and may be regenerated.", FONT),
        ("After editing: save this file, then run  python code/import_review.py labels/review/FedDrift_label_review.xlsx", FONT),
        ("", FONT),
        ("Sheets", BOLD),
        ("Judgment calls - 5 project-level decisions. 'decision' must be one of the option keys listed in 'options'.", FONT),
        ("Rules - 7 routine rules covering 3,667 routine events. decision: approve or reject.", FONT),
        ("Releases - 220 agency releases covering 442 non-routine events. decision: approve (use recommended cause), "
         "set (use your final_cause), or unknown (no reliable source).", FONT),
        ("Event overrides - optional single-event exceptions (event_id, final_cause, reason).", FONT),
        ("Events - read-only inventory of all 4,109 drift events with their rule or release.", FONT),
        ("Sources - read-only registry of every agency document cited (URL, publication date, SHA-256).", FONT),
        ("", FONT),
        ("Confidence grades", BOLD),
        ("strong = release-specific agency source and the observed footprint matches what the agency states", FONT),
        ("moderate = release-specific source without a numeric check, or program-level policy source consistent with the footprint", FONT),
        ("weak = timing/footprint only, or a located document that does not state the revision", FONT),
        ("none = no source located; recommended cause is 'unknown'", FONT),
        ("", FONT),
        ("Example of a completed Releases row (format only; not a real decision):", BOLD),
        ("decision = approve | final_cause = (leave blank to use recommended_cause) | verified_by = Tosin Clement | "
         "verified_date = 2026-10-05 | notes = checked Census 2026 annual revision introduction, p. 1", FONT),
        ("decision = set | final_cause = methodology_change | final_secondary_cause = seasonal_factor_recompute | "
         "verified_by = Tosin Clement | verified_date = 2026-10-05", FONT),
        ("Allowed final causes: " + ", ".join(final_causes), FONT),
    ]
    for text, font in lines:
        ws.append([text])
        ws.cell(ws.max_row, 1).font = font
    ws.column_dimensions["A"].width = 150

    w = wb.create_sheet("Judgment calls")
    style_sheet(w, jc, input_cols=("decision", "verified_by", "verified_date", "notes"),
                widths={"topic": 28, "evidence": 70, "options": 70, "recommendation": 18})
    w2 = wb.create_sheet("Rules")
    style_sheet(w2, rr, input_cols=("decision", "verified_by", "verified_date", "notes"),
                widths={"condition": 30, "rationale": 45, "agency_sources": 50, "agency_quotes": 70})
    add_list(w2, rr, "decision", ["approve", "reject"])
    w3 = wb.create_sheet("Releases")
    style_sheet(w3, rel, input_cols=("decision", "final_cause", "final_secondary_cause", "verified_by",
                                     "verified_date", "notes"),
                widths={"event_ids": 30, "sources": 55, "evidence_quote": 80, "automated_check": 55,
                        "recommendation_basis": 40, "series": 22})
    add_list(w3, rel, "decision", ["approve", "set", "unknown"])
    add_list(w3, rel, "final_cause", final_causes)
    add_list(w3, rel, "final_secondary_cause", final_causes)
    ov = pd.DataFrame(columns=["event_id", "final_cause", "final_secondary_cause", "reason", "source_url",
                               "source_ids", "verified_by", "verified_date", "notes"])
    w4 = wb.create_sheet("Event overrides")
    style_sheet(w4, pd.concat([ov, pd.DataFrame([[""] * len(ov.columns)] * 40, columns=ov.columns)]),
                input_cols=tuple(ov.columns), widths={"reason": 50, "source_url": 50})
    inv = ev[ev.is_drift_event][["event_id", "series_id", "vintage_prev", "vintage_next", "review_level",
                                 "proposed_rule", "release_cluster_id", "proposed_cause", "n_revised",
                                 "earliest_revised_obs", "revision_depth_months", "depth_censored",
                                 "mean_abs_pct_revision"]]
    w5 = wb.create_sheet("Events")
    style_sheet(w5, inv)
    w6 = wb.create_sheet("Sources")
    style_sheet(w6, reg, widths={"title": 60, "url": 70})
    import datetime as _dt
    wb.properties.created = wb.properties.modified = _dt.datetime(2026, 10, 3)   # deterministic file metadata
    wb.properties.creator = "FedDrift code/07_build_review.py"
    wb.save(os.path.join(REVIEW, "FedDrift_label_review.xlsx"))
    print(f"review package: {len(rr)} rules, {len(rel)} releases, {len(jc)} judgment calls, "
          f"{len(inv)} events, {len(reg)} sources")


if __name__ == "__main__":
    main()
